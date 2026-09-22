"""Shared helpers for building small in-memory DOCX/HTML/XML files for
extractor/endpoint tests — mirrors tests/pdf_helpers.py's approach so tests
stay fully offline with no binary fixture files committed to the repo.
"""

from __future__ import annotations

import io


def make_text_docx(
    heading: str = "Executive Summary",
    body: str = "Some body copy about the business.",
    band_fill: str | None = None,
) -> bytes:
    """A minimal .docx with a styled Heading 1 + a body paragraph, optionally
    with a shaded table cell (band-color detection)."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    doc = Document()
    heading_p = doc.add_paragraph(heading)
    heading_p.style = doc.styles["Heading 1"]
    heading_p.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    heading_p.runs[0].font.name = "Georgia"

    body_p = doc.add_paragraph(body)
    body_p.runs[0].font.name = "Arial"

    if band_fill:
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "Sample"
        tcPr = table.rows[0].cells[0]._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): band_fill})
        tcPr.append(shd)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_text_docx_with_image(
    heading: str = "Executive Summary",
    body: str = "Some body copy about the business.",
    image_size: tuple[int, int] = (80, 80),
    image_color: tuple[int, int, int] = (200, 50, 50),
) -> bytes:
    """Same as make_text_docx() plus one embedded PNG — exercises
    extract_reference_images()'s vision-attachment path."""
    from docx import Document
    from docx.shared import RGBColor
    from PIL import Image

    doc = Document()
    heading_p = doc.add_paragraph(heading)
    heading_p.style = doc.styles["Heading 1"]
    heading_p.runs[0].font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    heading_p.runs[0].font.name = "Georgia"

    body_p = doc.add_paragraph(body)
    body_p.runs[0].font.name = "Arial"

    img = Image.new("RGB", image_size, color=image_color)
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_buf.seek(0)
    doc.add_picture(img_buf)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_empty_docx() -> bytes:
    """A .docx with zero text runs — simulates a blank/unreadable document."""
    from docx import Document

    doc = Document()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_text_pptx(
    heading: str = "Executive Summary",
    body: str = "Some body copy about the business.",
    band_fill: str | None = None,
) -> bytes:
    """A minimal .pptx: one slide with a title placeholder + a content
    placeholder, optionally with a solid-filled decorative shape (band-color
    detection)."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # "Title and Content"

    title_run = slide.shapes.title.text_frame.paragraphs[0].add_run()
    title_run.text = heading
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(40)

    body_placeholder = slide.placeholders[1]
    body_run = body_placeholder.text_frame.paragraphs[0].add_run()
    body_run.text = body
    body_run.font.name = "Arial"

    if band_fill:
        from pptx.enum.shapes import MSO_SHAPE
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(1), Inches(1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(band_fill)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_text_pptx_with_image(
    heading: str = "Executive Summary",
    body: str = "Some body copy about the business.",
    image_size: tuple[int, int] = (80, 80),
    image_color: tuple[int, int, int] = (200, 50, 50),
) -> bytes:
    """Same as make_text_pptx() plus one embedded PNG — exercises
    extract_reference_images()'s vision-attachment path."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from PIL import Image

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])

    title_run = slide.shapes.title.text_frame.paragraphs[0].add_run()
    title_run.text = heading
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(40)

    body_placeholder = slide.placeholders[1]
    body_run = body_placeholder.text_frame.paragraphs[0].add_run()
    body_run.text = body
    body_run.font.name = "Arial"

    img = Image.new("RGB", image_size, color=image_color)
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_buf.seek(0)
    slide.shapes.add_picture(img_buf, Inches(1), Inches(1))

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_text_pptx_with_placeholder_image(
    image_size: tuple[int, int] = (80, 80),
    image_color: tuple[int, int, int] = (200, 50, 50),
) -> bytes:
    """A .pptx where the image is inserted into a picture PLACEHOLDER (the
    'Picture with Caption' layout), not via shapes.add_picture() — a very
    common real-world pattern (logo/photo placeholders on title-slide
    layouts) that reports shape_type == PLACEHOLDER rather than PICTURE,
    exercising extract_reference_images()'s placeholder-image path."""
    from pptx import Presentation
    from PIL import Image

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])  # "Picture with Caption"
    pic_placeholder = next(
        ph for ph in slide.placeholders
        if "PICTURE" in str(ph.placeholder_format.type)
    )
    img = Image.new("RGB", image_size, color=image_color)
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_buf.seek(0)
    pic_placeholder.insert_picture(img_buf)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pptx_with_table(
    heading: str = "Financial Information",
    cells: tuple[tuple[str, str], ...] = (("Revenue", "$1,000,000"), ("EBITDA", "$250,000")),
) -> bytes:
    """A .pptx with a title placeholder plus a table shape — a table's own
    shape has_text_frame=False (its cells each have their own text_frame),
    exercising the table-cell-text path in extract_plain_text() /
    extract_style_profile()'s run iteration."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_run = slide.shapes.title.text_frame.paragraphs[0].add_run()
    title_run.text = heading
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(40)

    gframe = slide.shapes.add_table(len(cells), 2, Inches(1), Inches(2), Inches(4), Inches(1.5))
    for row_idx, (label, value) in enumerate(cells):
        gframe.table.cell(row_idx, 0).text = label
        gframe.table.cell(row_idx, 1).text = value

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pptx_with_grouped_shapes(
    heading: str = "Executive Summary",
    band_fill: str = "2E5E3E",
    image_size: tuple[int, int] = (60, 60),
    image_color: tuple[int, int, int] = (9, 9, 9),
    grouped_text: str = "Grouped label text",
) -> bytes:
    """A .pptx with a title, plus a filled shape + a picture + a text box all
    grouped together — slide.shapes only lists TOP-LEVEL shapes, so a
    grouped logo/decorative-element cluster (a common, ordinary PowerPoint
    authoring pattern) is otherwise invisible to every extraction path."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    from PIL import Image

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_run = slide.shapes.title.text_frame.paragraphs[0].add_run()
    title_run.text = heading
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(40)

    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(3), Inches(1), Inches(1))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor.from_string(band_fill)

    img = Image.new("RGB", image_size, color=image_color)
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    img_buf.seek(0)
    pic = slide.shapes.add_picture(img_buf, Inches(2), Inches(3))

    tb = slide.shapes.add_textbox(Inches(3), Inches(3), Inches(2), Inches(1))
    tb.text_frame.text = grouped_text

    slide.shapes.add_group_shape([rect, pic, tb])

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_text_pptx_with_duplicate_logo(
    heading: str = "Executive Summary",
) -> bytes:
    """A .pptx where the SAME image is embedded on two different slides — a
    common real-world pattern (a company logo repeated across slides).
    Exercises extract_reference_images()'s dedup-by-content-hash path: the
    duplicate must not eat a second slot away from a genuinely different
    image."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from PIL import Image

    prs = Presentation()

    logo = Image.new("RGB", (50, 50), color=(10, 20, 30))
    logo_buf = io.BytesIO()
    logo.save(logo_buf, format="PNG")
    logo_bytes = logo_buf.getvalue()

    distinct = Image.new("RGB", (50, 50), color=(200, 100, 50))
    distinct_buf = io.BytesIO()
    distinct.save(distinct_buf, format="PNG")
    distinct_bytes = distinct_buf.getvalue()

    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    title_run = slide1.shapes.title.text_frame.paragraphs[0].add_run()
    title_run.text = heading
    title_run.font.size = Pt(40)
    slide1.shapes.add_picture(io.BytesIO(logo_bytes), Inches(0), Inches(0))

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.shapes.add_picture(io.BytesIO(logo_bytes), Inches(0), Inches(0))       # same logo again
    slide2.shapes.add_picture(io.BytesIO(distinct_bytes), Inches(2), Inches(2))   # genuinely different image

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_empty_pptx() -> bytes:
    """A .pptx with zero slides — simulates a blank/unreadable deck."""
    from pptx import Presentation

    prs = Presentation()
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_text_html(
    heading: str = "Executive Summary",
    body: str = "Some body copy about the business.",
) -> bytes:
    """A minimal HTML page with a <style> block defining heading/body/band colors."""
    html = f"""<html><head><style>
h1 {{ color: #8B0000; font-family: Georgia, serif; text-align: center; }}
.band {{ background-color: #F4E4C1; }}
p {{ color: #333333; font-family: Arial, sans-serif; }}
</style></head><body>
<h1 class="band">{heading}</h1>
<p>{body}</p>
</body></html>"""
    return html.encode("utf-8")


def make_empty_html() -> bytes:
    return b"<html><head></head><body></body></html>"


def make_custom_schema_xml() -> bytes:
    """A hand-rolled XML schema with real text in non-HTML tags — no h1/p/etc,
    exercises the flat-text fallback path (no design hints found, degrades to
    defaults rather than erroring on a genuinely non-empty file)."""
    return (
        b"<template><section name='Executive Summary'>"
        b"Some body copy about the business.</section></template>"
    )
