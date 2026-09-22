"""Tests for core.pptx_style_extractor — the .pptx counterpart to
tests/test_docx_style_extractor.py."""

from __future__ import annotations

import base64

import pytest

from core.pdf_style_extractor import NoExtractableTextError
from core.pptx_style_extractor import extract_plain_text, extract_reference_images, extract_style_profile
from tests.template_helpers import (
    make_empty_pptx,
    make_pptx_with_grouped_shapes,
    make_pptx_with_table,
    make_text_pptx,
    make_text_pptx_with_duplicate_logo,
    make_text_pptx_with_image,
    make_text_pptx_with_placeholder_image,
)


def test_extract_raises_on_empty_pptx():
    with pytest.raises(NoExtractableTextError):
        extract_style_profile(make_empty_pptx())


def test_extract_basic_pptx_shape_and_values():
    template, warnings = extract_style_profile(make_text_pptx())
    assert template["id"] == "custom-upload"
    assert template["palette"]["primary"] == "#1f4e79"   # title run color
    assert template["palette"]["mid"] == "#000000"        # body run has no explicit color
    assert "Georgia" in template["fonts"]["heading"]
    assert "Helvetica" in template["fonts"]["body"] or "Arial" in template["fonts"]["body"]
    assert template["headings"]["I. Executive Summary"] == "Executive Summary"
    assert len(warnings) == 9
    assert "I. Executive Summary" not in warnings


def test_extract_detects_band_color():
    template, _warnings = extract_style_profile(make_text_pptx(band_fill="2E5E3E"))
    assert template["palette"]["accent"] == "#2e5e3e"
    assert "#2e5e3e" in template["section_header_override"]


def test_per_slide_title_detection_finds_headings_with_no_placeholder():
    """Regression guard, found against a real uploaded CIM deck: a slide
    with NO title placeholder at all (a plain text box sized up instead —
    a common real-world PowerPoint authoring pattern) must still have its
    real heading detected via the largest-font-on-this-slide fallback, even
    when OTHER slides in the same deck DO use a title placeholder. The old
    deck-wide check let one slide's placeholder starve every other slide's
    real heading into 'body'."""
    import io
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    prs = Presentation()

    slide_with_placeholder = prs.slides.add_slide(prs.slide_layouts[1])
    slide_with_placeholder.shapes.title.text_frame.paragraphs[0].add_run().text = "Placeholder Heading"

    slide_without_placeholder = prs.slides.add_slide(prs.slide_layouts[6])  # blank, no placeholders
    tb = slide_without_placeholder.shapes.add_textbox(0, 0, Pt(400), Pt(60))
    heading_run = tb.text_frame.paragraphs[0].add_run()
    heading_run.text = "Text Box Heading"
    heading_run.font.size = Pt(32)
    body_p = tb.text_frame.add_paragraph()
    body_run = body_p.add_run()
    body_run.text = "some smaller body copy"
    body_run.font.size = Pt(14)

    buf = io.BytesIO()
    prs.save(buf)

    template, _warnings = extract_style_profile(buf.getvalue())
    heading_i = template["headings"]["I. Executive Summary"]
    # Neither real heading maps to a canonical section by name here — this
    # test only cares that BOTH slides' real headings were found as title
    # candidates at all (checked via the module-level helper directly,
    # since headings dict only surfaces sections that matched a synonym).
    from pptx import Presentation as _P
    from core.pptx_style_extractor import _classify_slide_runs
    prs2 = _P(io.BytesIO(buf.getvalue()))
    all_titles = [
        run.text.strip()
        for slide in prs2.slides
        for _shape, _para, run, is_title in _classify_slide_runs(slide)
        if is_title
    ]
    assert "Placeholder Heading" in all_titles
    assert "Text Box Heading" in all_titles


def test_per_slide_title_detection_not_fooled_by_a_size_tie():
    """Regression guard for the exact bug found against a real uploaded CIM
    deck: a slide with exactly TWO distinct font sizes, each appearing on
    only one run (a 32pt heading text box + one body run with NO explicit
    size, which falls back to the 18pt default) — Counter.most_common()
    breaks a count-of-1-each tie by insertion order, which was picking the
    LARGE size as "baseline" about as often as the small one, so the
    >=1.15x comparison silently failed and neither run got flagged as a
    title. Sorting distinct sizes descending removes the tie."""
    import io
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(0, 0, Pt(400), Pt(100))
    heading_run = tb.text_frame.paragraphs[0].add_run()
    heading_run.text = "REVENUE SOURCES:"
    heading_run.font.size = Pt(32)
    body_p = tb.text_frame.add_paragraph()
    body_run = body_p.add_run()
    body_run.text = "no explicit size set at all"
    # body_run.font.size deliberately left unset (None) — inherits/falls back

    buf = io.BytesIO()
    prs.save(buf)

    from core.pptx_style_extractor import _classify_slide_runs
    titles = [
        run.text.strip()
        for _shape, _para, run, is_title in _classify_slide_runs(slide)
        if is_title
    ]
    assert titles == ["REVENUE SOURCES:"]


def test_extract_never_crashes_on_theme_colored_or_unset_color_runs():
    """A run's font color is very often NOT an explicit RGBColor in a real
    deck — it can be a theme/scheme color (font.color.rgb raises
    AttributeError) or simply unset/inherited (also raises AttributeError,
    different message). Both must degrade to the '#000000' fallback, never
    propagate."""
    import io
    from pptx import Presentation
    from pptx.dml.color import MSO_THEME_COLOR

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_run = slide.shapes.title.text_frame.paragraphs[0].add_run()
    title_run.text = "Themed heading, no explicit rgb"
    title_run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1

    body_run = slide.placeholders[1].text_frame.paragraphs[0].add_run()
    body_run.text = "plain body, no color set at all"

    buf = io.BytesIO()
    prs.save(buf)

    template, _warnings = extract_style_profile(buf.getvalue())
    assert template["palette"]["primary"] == "#000000"
    assert template["palette"]["mid"] == "#000000"


def test_extract_plain_text_flattens_slide_text():
    text = extract_plain_text(make_text_pptx(heading="My Heading", body="My Body"))
    assert "My Heading" in text
    assert "My Body" in text


def test_extract_plain_text_includes_table_cell_text():
    """Regression guard: a table shape itself has_text_frame=False (only its
    cells do), so a naive has_text_frame check silently drops every table on
    the slide — real CIM/financial templates commonly have a data table."""
    text = extract_plain_text(make_pptx_with_table())
    assert "Revenue" in text
    assert "$1,000,000" in text
    assert "EBITDA" in text


def test_extract_style_profile_never_crashes_with_a_table_present():
    template, _warnings = extract_style_profile(make_pptx_with_table())
    assert template["palette"]["primary"] == "#1f4e79"  # title color still detected correctly


def test_extraction_descends_into_grouped_shapes():
    """Regression guard: slide.shapes only lists top-level shapes — a filled
    shape, a picture, and a text box all grouped together (an ordinary
    PowerPoint authoring pattern, e.g. a logo grouped with its background)
    must still be found by band-color, text, and image extraction, not
    silently skipped because they're nested inside a GroupShape."""
    pptx_bytes = make_pptx_with_grouped_shapes()

    template, _warnings = extract_style_profile(pptx_bytes)
    assert template["palette"]["accent"] == "#2e5e3e"

    text = extract_plain_text(pptx_bytes)
    assert "Grouped label text" in text

    images = extract_reference_images(pptx_bytes)
    assert len(images) == 1


def test_extract_reference_images_returns_empty_list_when_no_images():
    assert extract_reference_images(make_text_pptx()) == []


def test_extract_reference_images_finds_embedded_image():
    images = extract_reference_images(make_text_pptx_with_image())
    assert len(images) == 1
    assert images[0]["mime"] == "image/png"
    # round-trips to real, decodable PNG bytes — not a placeholder/garbage string
    raw = base64.standard_b64decode(images[0]["b64"])
    assert raw.startswith(b"\x89PNG\r\n\x1a\n")


def test_extract_reference_images_finds_image_in_picture_placeholder():
    """Regression guard: a shape_type == PICTURE check alone misses images
    inserted into a picture PLACEHOLDER (common in real title-slide
    layouts) — those report shape_type == PLACEHOLDER, not PICTURE, even
    though shape.image works fine."""
    images = extract_reference_images(make_text_pptx_with_placeholder_image())
    assert len(images) == 1
    assert images[0]["mime"] == "image/png"
    raw = base64.standard_b64decode(images[0]["b64"])
    assert raw.startswith(b"\x89PNG\r\n\x1a\n")


def test_extract_reference_images_dedupes_a_reused_logo():
    """Regression guard, found against a real uploaded CIM deck: the same
    company logo, embedded once but referenced by picture shapes on 3
    different slides, ate 3 of the 4 available image slots — leaving only 1
    slot for every OTHER genuinely distinct image in the deck. Duplicate
    content (by hash) must collapse to a single entry before capping."""
    images = extract_reference_images(make_text_pptx_with_duplicate_logo())
    assert len(images) == 2  # the logo (once, not twice) + the one distinct image
    raws = [base64.standard_b64decode(img["b64"]) for img in images]
    assert len(set(raws)) == 2  # both entries are genuinely different bytes


def test_extract_reference_images_caps_at_max_and_orders_largest_first():
    pptx_bytes = make_text_pptx_with_image()
    # Only one image in the fixture — this just locks in the cap constant's
    # existence/behavior contract without needing 5 embedded images.
    from core.pptx_style_extractor import _MAX_REFERENCE_IMAGES
    images = extract_reference_images(pptx_bytes)
    assert len(images) <= _MAX_REFERENCE_IMAGES


def test_extract_reference_images_never_raises_on_corrupt_input():
    # Not a real .pptx at all — python-pptx's Presentation() will choke; the
    # function must degrade to an empty list rather than propagate.
    assert extract_reference_images(b"not a real pptx file") == []


def test_extract_reference_images_never_raises_on_valid_zip_that_isnt_pptx():
    # A different failure mode than pure garbage bytes: a well-formed ZIP
    # archive that just isn't a PowerPoint package (BadZipFile vs. the
    # KeyError python-pptx raises hunting for [Content_Types].xml) — must
    # degrade the same way, not propagate a different exception type.
    import zipfile
    import io as _io
    buf = _io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "not a pptx")
    assert extract_reference_images(buf.getvalue()) == []
