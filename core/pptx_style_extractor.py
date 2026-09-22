"""Deterministic PPTX -> CIM template-style extractor. No LLM involved.

PowerPoint counterpart to core/docx_style_extractor.py — same output contract
(see core/pdf_style_extractor.py's docstring): a plain dict shaped like one
entry of core/templates.py's TEMPLATES, plus a list of canonical sections
that had no matching heading. Different source library (python-pptx instead
of python-docx/PyMuPDF) and a slide-based object model instead of a
paragraph or page/bbox one — heading vs. body is detected PER SLIDE: a
slide's title placeholder (TITLE/CENTER_TITLE) wins when present, otherwise
that slide's own single largest-font run becomes its title (mirroring the
PDF/DOCX extractors' largest-font heuristic, scoped to one slide rather than
the whole deck — see _classify_slide_runs for why deck-wide isn't good
enough for real templates).

Legacy binary .ppt is NOT parseable by python-pptx — Presentation() raises,
and the caller (core/template_extractor.py) turns that into a clear
user-facing error rather than a 500.
"""

from __future__ import annotations

import base64
import hashlib
import io
from collections import Counter

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN

from core.pdf_style_extractor import NoExtractableTextError
from core.section_matcher import CANONICAL_SECTIONS, match_sections
from core.style_shared import detect_bullet_glyph, font_stack_for

# SUBTITLE deliberately excluded — it's slide-deck tagline/body copy, not a
# heading, and pulling it into section-heading matching would give wrong matches.
_TITLE_PLACEHOLDER_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)


def _iter_all_shapes(shapes):
    """Recursively yield every shape, descending into grouped shapes.
    slide.shapes only lists TOP-LEVEL shapes — a logo grouped with its
    background, or any grouped decorative element (a common, ordinary
    PowerPoint authoring pattern), sits nested inside a GroupShape's own
    .shapes collection and is otherwise completely invisible to every scan
    in this module (band color, text runs, embedded images)."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_all_shapes(shape.shapes)
        else:
            yield shape


def _is_title_shape(shape) -> bool:
    if not shape.is_placeholder:
        return False
    try:
        return shape.placeholder_format.type in _TITLE_PLACEHOLDER_TYPES
    except (AttributeError, ValueError):
        return False


def _hex_from_run(run) -> str | None:
    """A run's font color may be unset (inherits from theme/layout), or set
    via a theme color rather than an explicit RGB value — .rgb raises
    AttributeError in that case. Either way there's no explicit hex to read."""
    try:
        rgb = run.font.color.rgb
    except (AttributeError, TypeError):
        return None
    return f"#{str(rgb).lower()}" if rgb is not None else None


def _pt_size(run) -> float:
    size = run.font.size
    return size.pt if size is not None else 18.0  # python-pptx has no single universal default; 18pt is a reasonable body-text fallback


def _classify_slide_runs(slide):
    """Yield (shape, paragraph, run, is_title) for every non-empty text run
    on ONE slide, deciding title vs. body PER SLIDE rather than deck-wide.

    A run is a title if it sits in a TITLE/CENTER_TITLE placeholder. If the
    slide has no such placeholder AT ALL, the slide's own single
    largest-font run becomes its title instead — provided that size is
    meaningfully bigger (>=1.15x) than the rest of the slide's text, mirroring
    the PDF/DOCX extractors' largest-font heuristic but scoped to one slide.
    Table-cell text is always body — financial/comparison tables are content,
    never headings.

    Why per-slide, not deck-wide: real PowerPoint templates very often use a
    native title placeholder on some slides (e.g. the cover) and plain,
    custom-positioned text boxes for section headers on others — confirmed
    against a real uploaded CIM template where 2 of its 4 actual section
    headings were sized-up plain text boxes on slides with no placeholder at
    all. A deck-wide "does title_runs exist anywhere" check let one
    incidental slide's placeholder (e.g. a testimonial-quote slide that
    happens to reuse the CENTER_TITLE placeholder) starve every other
    slide's real, non-placeholder heading into 'body' — only 1 of that
    deck's ~4 real section headings was ever detected.
    """
    placeholder_title: list[tuple] = []
    text_candidates: list[tuple] = []
    table_body: list[tuple] = []

    for shape in _iter_all_shapes(slide.shapes):
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            if run.text.strip():
                                table_body.append((shape, para, run))
            continue
        if not shape.has_text_frame:
            continue
        is_ph_title = _is_title_shape(shape)
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if not run.text.strip():
                    continue
                (placeholder_title if is_ph_title else text_candidates).append((shape, para, run))

    if placeholder_title:
        title, body = placeholder_title, text_candidates
    elif text_candidates:
        # Compare the largest distinct size against the SECOND-largest distinct
        # size — not a "most common size" mode. With few runs on a slide, two
        # different sizes very often each occur exactly once; Counter.most_common()
        # breaks that tie by insertion order, which silently picked the LARGE
        # size as "baseline" about as often as the small one, defeating the
        # >=1.15x comparison outright (a real slide with one 32pt heading run
        # and one body run at an inherited/unset size classified neither as a
        # title). Sorting distinct sizes removes the tie entirely.
        sized = [(round(_pt_size(t[2])), t) for t in text_candidates]
        distinct_desc = sorted({sz for sz, _ in sized}, reverse=True)
        if len(distinct_desc) > 1 and distinct_desc[0] >= distinct_desc[1] * 1.15:
            top = distinct_desc[0]
            title = [t for sz, t in sized if sz == top]
            body = [t for sz, t in sized if sz != top]
        else:
            title, body = [], text_candidates
    else:
        title, body = [], []

    for shape, para, run in title:
        yield shape, para, run, True
    for shape, para, run in body + table_body:
        yield shape, para, run, False


def _iter_text_runs(prs: "Presentation"):
    """Yield (shape, paragraph, run, is_title) for every non-empty text run
    across every slide, in slide order — see _classify_slide_runs for how
    title vs. body is decided."""
    for slide in prs.slides:
        yield from _classify_slide_runs(slide)


def _colors(runs: list[tuple]) -> Counter:
    hexes = [h for h in (_hex_from_run(r) for _, r in runs) if h]
    return Counter(hexes) if hexes else Counter(["#000000"])


def _band_color(prs: "Presentation") -> str | None:
    """Most common non-white/black solid shape fill across every slide, if
    any — the PPTX analog of the PDF extractor's vector-fill scan and the
    DOCX extractor's table-cell-shading scan (decorative color bands/blocks
    are commonly drawn as plain filled shapes in slide decks)."""
    fills: Counter[str] = Counter()
    for slide in prs.slides:
        for shape in _iter_all_shapes(slide.shapes):
            try:
                if shape.fill.type != MSO_FILL_TYPE.SOLID:
                    continue
                rgb = shape.fill.fore_color.rgb
            except (AttributeError, TypeError, ValueError):
                continue
            if rgb is None:
                continue
            hex_color = f"#{str(rgb).lower()}"
            if hex_color not in ("#ffffff", "#000000"):
                fills[hex_color] += 1
    return fills.most_common(1)[0][0] if fills else None


def extract_plain_text(pptx_bytes: bytes) -> str:
    """Flatten a .pptx to plain text (every text shape's text, plus every
    table's cells tab-joined per row, slide by slide). Used only to give
    Claude a readable copy of an uploaded template file when there's no
    vision path for it — see core/llm.py:generate_cim_html, where Claude's
    document content block only accepts application/pdf.
    """
    prs = Presentation(io.BytesIO(pptx_bytes))
    lines: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = []
        for shape in _iter_all_shapes(slide.shapes):
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = "\t".join(cell.text_frame.text.strip() for cell in row.cells)
                    if row_text.strip():
                        slide_lines.append(row_text)
            elif shape.has_text_frame and shape.text_frame.text.strip():
                slide_lines.append(shape.text_frame.text.strip())
        if slide_lines:
            lines.append(f"--- Slide {i} ---")
            lines.extend(slide_lines)
    return "\n".join(lines)


_MAX_REFERENCE_IMAGES = 4
_ACCEPTED_IMAGE_CONTENT_TYPES = {"image/png", "image/jpeg", "image/webp", "image/gif"}


def extract_reference_images(pptx_bytes: bytes) -> list[dict]:
    """Pull embedded images out of a .pptx package (logos, photos,
    decorative graphics) so Claude can actually SEE the template's visual
    design instead of relying solely on the flattened text from
    extract_plain_text() above. Returns up to _MAX_REFERENCE_IMAGES dicts:
    {"b64": str, "mime": str}, largest-by-byte-size first among DISTINCT
    images — a company logo embedded once in the package but referenced by
    multiple picture shapes across slides (a very common real-world pattern:
    confirmed against a real uploaded CIM deck where the same logo appeared
    on 3 different slides) is deduplicated by content hash first, so the cap
    goes to genuinely different images instead of 3 copies of the same one
    crowding out the deck's other real photos/screenshots. Best-effort —
    truly never raises: an unreadable image shape is skipped, and even
    pptx_bytes that isn't a valid .pptx at all degrades to an empty list
    rather than propagating (this is a design-reference nice-to-have, never
    worth failing the whole template upload over).
    """
    found: list[tuple[int, dict]] = []
    seen_hashes: set[bytes] = set()
    try:
        prs = Presentation(io.BytesIO(pptx_bytes))
        for slide in prs.slides:
            for shape in _iter_all_shapes(slide.shapes):
                # Deliberately NOT gated on shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                # a picture inserted into a picture PLACEHOLDER (very common in real
                # title-slide layouts, e.g. a logo/photo placeholder) reports
                # shape_type == PLACEHOLDER, not PICTURE, even though it has a real
                # .image — that gate silently skipped every placeholder-inserted
                # image. Instead just try .image on every shape; a shape that
                # doesn't have one (text box, autoshape, table, chart, an empty
                # placeholder) raises a plain AttributeError, caught below.
                try:
                    image = shape.image
                    content_type = (image.content_type or "").split(";")[0].strip().lower()
                    if content_type not in _ACCEPTED_IMAGE_CONTENT_TYPES:
                        continue
                    blob = image.blob
                    if not blob:
                        continue
                    blob_hash = hashlib.md5(blob).digest()
                    if blob_hash in seen_hashes:
                        continue
                    seen_hashes.add(blob_hash)
                    b64 = base64.standard_b64encode(blob).decode("ascii")
                    found.append((len(blob), {"b64": b64, "mime": content_type}))
                except Exception:
                    continue  # one unreadable image shape must never fail the whole upload
    except Exception:
        return []  # pptx_bytes wasn't a valid .pptx package at all

    found.sort(key=lambda pair: pair[0], reverse=True)
    return [img for _size, img in found[:_MAX_REFERENCE_IMAGES]]


def extract_style_profile(pptx_bytes: bytes) -> tuple[dict, list[str]]:
    """Extract a CIM template-style dict from an uploaded .pptx. Deterministic, no LLM.

    Same return contract as core/pdf_style_extractor.py:extract_style_profile.
    Raises NoExtractableTextError if the deck has no text runs at all.
    """
    prs = Presentation(io.BytesIO(pptx_bytes))
    all_runs = list(_iter_text_runs(prs))
    runs = [(shape, run) for shape, _para, run, _is_title in all_runs]
    title_runs = [(shape, run) for shape, _para, run, is_title in all_runs if is_title]
    body_runs = [(shape, run) for shape, _para, run, is_title in all_runs if not is_title]

    if not runs:
        raise NoExtractableTextError(
            "This PowerPoint file has no readable text — we can't read its layout."
        )

    if not title_runs:
        # Pathological last resort: _classify_slide_runs already applies a
        # per-slide placeholder-or-largest-font rule, so this only fires if
        # literally every slide has neither a title placeholder nor any
        # size variation at all in its text.
        title_runs, body_runs = runs[:1], runs[1:]

    heading_colors = _colors(title_runs)
    body_colors = _colors(body_runs) if body_runs else heading_colors

    primary = heading_colors.most_common(1)[0][0]
    mid = body_colors.most_common(1)[0][0]

    band = _band_color(prs)
    accent = band or (mid if mid != primary else "#3a6ea5")
    light = "#ffffff"

    top_title_shape, top_title_run = title_runs[0]
    top_body_shape, top_body_run = body_runs[0] if body_runs else title_runs[0]
    fonts = {
        "heading": font_stack_for(top_title_run.font.name or ""),
        "body": font_stack_for(top_body_run.font.name or ""),
    }

    bullet_glyph = detect_bullet_glyph([r.text for _, r in body_runs])

    top_title_para = next(
        (p for p in top_title_shape.text_frame.paragraphs if any(r.text.strip() for r in p.runs)),
        None,
    )
    alignment = "center" if top_title_para is not None and top_title_para.alignment == PP_ALIGN.CENTER else "left"

    title_texts = [r.text.strip() for _, r in title_runs if r.text.strip()]
    matched = match_sections(title_texts)
    headings = {h: h for h in CANONICAL_SECTIONS}
    headings.update(matched)
    warnings = [h for h in CANONICAL_SECTIONS if h not in matched]

    band_sentence = (
        f"Section headers sit on a solid {band} background band."
        if band else
        "Section headers use a plain background with a thin rule beneath the title."
    )
    layout_notes = (
        f"{band_sentence} Bullet lists use the '{bullet_glyph}' glyph as the marker, "
        "matching the uploaded template."
    )
    align_sentence = (
        "Cover content is centered, matching the uploaded template's title slide."
        if alignment == "center" else
        "Cover content is left-aligned, matching the uploaded template's title slide."
    )
    cover_override = align_sentence + (f" The title sits on a {band} background band." if band else "")
    section_header_override = (
        f"Section header band uses {band} as the background color, matching the uploaded template."
        if band else
        "Section headers use a plain background with a rule beneath the title, matching the uploaded template."
    )

    template = {
        "id": "custom-upload",
        "name": "Your Uploaded Template",
        "tagline": "Matches the design of your uploaded PowerPoint template.",
        "allow_brand_override": False,
        "palette": {"primary": primary, "accent": accent, "light": light, "mid": mid},
        "fonts": fonts,
        "layout_notes": layout_notes,
        "cover_override": cover_override,
        "section_header_override": section_header_override,
        "headings": headings,
    }
    return template, warnings
