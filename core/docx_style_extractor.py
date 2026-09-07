"""Deterministic DOCX -> CIM template-style extractor. No LLM involved.

Word-document counterpart to core/pdf_style_extractor.py — same output
contract (see that module's docstring): a plain dict shaped like one entry
of core/templates.py's TEMPLATES, plus a list of canonical sections that had
no matching heading. Different source library (python-docx instead of
PyMuPDF), because .docx has no page/bbox geometry — only a paragraph/run
object model. Heading vs. body is detected primarily via Word's own
"Heading"/"Title" paragraph styles (how real templates are authored), and
only falls back to a largest-font-size heuristic (mirroring the PDF
extractor) when no such styles are used at all.

Legacy binary .doc is NOT parseable by python-docx — Document() raises, and
the caller (core/template_extractor.py) turns that into a clear user-facing
error rather than a 500.
"""

from __future__ import annotations

import io
from collections import Counter

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from core.pdf_style_extractor import NoExtractableTextError
from core.section_matcher import CANONICAL_SECTIONS, match_sections
from core.style_shared import detect_bullet_glyph, font_stack_for

_HEADING_STYLE_PREFIXES = ("heading", "title")


def _hex_from_rgb(rgb) -> str | None:
    """python-docx RGBColor stringifies to 'RRGGBB' (no '#', uppercase) — or the
    run has no explicit color at all (theme/inherited), in which case rgb is None."""
    return f"#{str(rgb).lower()}" if rgb is not None else None


def _colors(runs: list[tuple]) -> Counter:
    hexes = [h for h in (_hex_from_rgb(r.font.color.rgb) for _, r in runs) if h]
    return Counter(hexes) if hexes else Counter(["#000000"])


def _iter_runs(doc: "Document"):
    """Yield (paragraph, run) for every non-empty run in the document body,
    including runs inside table cells."""
    for para in doc.paragraphs:
        for run in para.runs:
            if run.text.strip():
                yield para, run
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.text.strip():
                            yield para, run


def _is_heading_para(para) -> bool:
    style_name = (para.style.name if para.style else "") or ""
    return style_name.strip().lower().startswith(_HEADING_STYLE_PREFIXES)


def _band_color(doc: "Document") -> str | None:
    """Most common non-white/black table-cell shading fill across the document,
    if any — the DOCX analog of core/pdf_style_extractor.py's drawing-fill scan."""
    fills: Counter[str] = Counter()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.tcPr
                if tcPr is None:
                    continue
                shd = tcPr.find(qn("w:shd"))
                if shd is None:
                    continue
                fill = shd.get(qn("w:fill"))
                if fill and fill.upper() not in ("FFFFFF", "000000", "AUTO"):
                    fills[f"#{fill.lower()}"] += 1
    return fills.most_common(1)[0][0] if fills else None


def extract_plain_text(docx_bytes: bytes) -> str:
    """Flatten a .docx to plain text (paragraphs + table cells, tab-joined).

    Used only to give Claude a readable copy of an uploaded template file when
    there's no vision path for it — see core/llm.py:generate_cim_html, where
    Claude's document content block only accepts application/pdf.
    """
    doc = Document(io.BytesIO(docx_bytes))
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def extract_style_profile(docx_bytes: bytes) -> tuple[dict, list[str]]:
    """Extract a CIM template-style dict from an uploaded .docx. Deterministic, no LLM.

    Same return contract as core/pdf_style_extractor.py:extract_style_profile.
    Raises NoExtractableTextError if the document has no text runs at all.
    """
    doc = Document(io.BytesIO(docx_bytes))
    runs = list(_iter_runs(doc))
    if not runs:
        raise NoExtractableTextError(
            "This Word document has no readable text — we can't read its layout."
        )

    heading_runs = [(p, r) for p, r in runs if _is_heading_para(p)]
    body_runs = [(p, r) for p, r in runs if not _is_heading_para(p)]

    if not heading_runs:
        # No named Heading/Title styles used — fall back to the largest-font-size
        # heuristic, mirroring the PDF extractor's approach. A run with no explicit
        # size inherits its style's size, which we can't resolve here, so it's
        # treated as Word's own default body size (11pt).
        def _pt(run) -> float:
            size = run.font.size
            return size.pt if size is not None else 11.0

        body_size = Counter(round(_pt(r)) for _, r in runs).most_common(1)[0][0]
        heading_runs = [(p, r) for p, r in runs if _pt(r) >= body_size * 1.15]
        body_runs = [(p, r) for p, r in runs if _pt(r) < body_size * 1.15]
        if not heading_runs:
            heading_runs, body_runs = runs[:1], runs[1:]

    heading_colors = _colors(heading_runs)
    body_colors = _colors(body_runs) if body_runs else heading_colors

    primary = heading_colors.most_common(1)[0][0]
    mid = body_colors.most_common(1)[0][0]

    band = _band_color(doc)
    accent = band or (mid if mid != primary else "#3a6ea5")
    light = "#ffffff"

    _, top_heading_run = heading_runs[0]
    _, top_body_run = body_runs[0] if body_runs else heading_runs[0]
    fonts = {
        "heading": font_stack_for(top_heading_run.font.name or ""),
        "body": font_stack_for(top_body_run.font.name or ""),
    }

    bullet_glyph = detect_bullet_glyph([p.text for p, _ in body_runs])

    top_para, _ = heading_runs[0]
    alignment = "center" if top_para.alignment == WD_ALIGN_PARAGRAPH.CENTER else "left"

    matched = match_sections([p.text.strip() for p, _ in heading_runs])
    headings = {h: h for h in CANONICAL_SECTIONS}
    headings.update(matched)
    warnings = [h for h in CANONICAL_SECTIONS if h not in matched]

    band_sentence = (
        f"Section headers sit on a solid {band} background band."
        if band else
        "Section headers use a plain background with a thin rule beneath the title."
    )
    layout_notes = (
        f"{band_sentence} Bullet lists use the '{bullet_glyph}' glyph as the marker, "
        "matching the uploaded template."
    )
    align_sentence = (
        "Cover content is centered, matching the uploaded template's title page."
        if alignment == "center" else
        "Cover content is left-aligned, matching the uploaded template's title page."
    )
    cover_override = align_sentence + (f" The title sits on a {band} background band." if band else "")
    section_header_override = (
        f"Section header band uses {band} as the background color, matching the uploaded template."
        if band else
        "Section headers use a plain background with a rule beneath the title, matching the uploaded template."
    )

    template = {
        "id": "custom-upload",
        "name": "Your Uploaded Template",
        "tagline": "Matches the design of your uploaded Word template.",
        "allow_brand_override": False,
        "palette": {"primary": primary, "accent": accent, "light": light, "mid": mid},
        "fonts": fonts,
        "layout_notes": layout_notes,
        "cover_override": cover_override,
        "section_header_override": section_header_override,
        "headings": headings,
    }
    return template, warnings
